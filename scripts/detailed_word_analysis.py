import docx
import json
import re

# Word dosyasını oku
doc = docx.Document('Sonuclar-Ortaokul-Rapor.docx')

print(f'Word dosyası toplam paragraf sayısı: {len(doc.paragraphs)}')

# Tüm paragrafları incele
print('\n=== İLK 50 PARAGRAF ===')
for i, paragraph in enumerate(doc.paragraphs[:50]):
    if paragraph.text.strip():
        print(f'{i+1:3d}. {paragraph.text.strip()}')

# Tabloları kontrol et
print(f'\nWord dosyasında tablo sayısı: {len(doc.tables)}')

if doc.tables:
    print('\n=== İLK TABLO ===')
    table = doc.tables[0]
    for i, row in enumerate(table.rows[:10]):  # İlk 10 satır
        row_data = []
        for cell in row.cells:
            row_data.append(cell.text.strip())
        print(f'Satır {i+1}: {row_data}')

# Tüm metni birleştir ve farklı aramalar yap
full_text = ''
for paragraph in doc.paragraphs:
    full_text += paragraph.text + '\n'

# Farklı anahtar kelimeler ara
keywords = ['Öğrenci', 'Adı', 'Soyadı', 'Cinsiyeti', 'Sınıfı', 'Öğrenme', 'Stili', 'AYRIŞTIRAN', 'YERLEŞTİREN', 'DEĞİŞTİREN', 'ÖZÜMSEYEN']

print('\n=== ANAHTAR KELİME ARAMALARI ===')
for keyword in keywords:
    matches = re.findall(rf'{keyword}[^\n]*', full_text, re.IGNORECASE)
    if matches:
        print(f'{keyword}: {len(matches)} eşleşme')
        for match in matches[:3]:  # İlk 3 eşleşme
            print(f'  - {match}')

# Metinde "8" geçen yerleri ara (sınıf bilgisi için)
grade_matches = re.findall(r'8[^\n]*', full_text)
print(f'\n"8" geçen yerler: {len(grade_matches)}')
for i, match in enumerate(grade_matches[:10]):
    print(f'{i+1}. {match}')

# Metinde isim benzeri kelimeler ara
name_pattern = r'[A-ZÇĞIİÖŞÜ][a-zçğıiöşü]+\s+[A-ZÇĞIİÖŞÜ][a-zçğıiöşü]+'
name_matches = re.findall(name_pattern, full_text)
print(f'\nİsim benzeri kelimeler: {len(name_matches)}')
for i, name in enumerate(name_matches[:20]):
    print(f'{i+1}. {name}')
