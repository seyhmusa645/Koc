import docx
import json

# Word dosyasını oku
doc = docx.Document('Sonuclar-Ortaokul-Rapor.docx')

print(f'Word dosyasında tablo sayısı: {len(doc.tables)}')

# İlk 10 tabloyu detaylı incele
for table_num in range(min(10, len(doc.tables))):
    table = doc.tables[table_num]
    print(f'\n=== TABLO {table_num + 1} ===')
    print(f'Satır sayısı: {len(table.rows)}')
    print(f'Sütun sayısı: {len(table.columns)}')
    
    # İlk 5 satırı göster
    for row_num, row in enumerate(table.rows[:5]):
        row_data = []
        for cell in row.cells:
            cell_text = cell.text.strip()
            if cell_text:  # Boş olmayan hücreleri göster
                row_data.append(cell_text)
        if row_data:
            print(f'Satır {row_num + 1}: {row_data}')

# Belirli bir tabloyu daha detaylı incele (örneğin 5. tablo)
if len(doc.tables) > 4:
    print(f'\n=== TABLO 5 DETAYLI ANALİZ ===')
    table = doc.tables[4]
    for row_num, row in enumerate(table.rows):
        print(f'\nSatır {row_num + 1}:')
        for col_num, cell in enumerate(row.cells):
            cell_text = cell.text.strip()
            if cell_text:
                print(f'  Sütun {col_num + 1}: "{cell_text}"')

# Tüm tablolarda "Öğrenci" kelimesini ara
print(f'\n=== "ÖĞRENCİ" KELİMESİ ARAMASI ===')
student_found = 0
for table_num, table in enumerate(doc.tables):
    for row in table.rows:
        for cell in row.cells:
            if 'Öğrenci' in cell.text:
                student_found += 1
                if student_found <= 5:  # İlk 5 eşleşmeyi göster
                    print(f'Tablo {table_num + 1}: {cell.text.strip()}')

print(f'\nToplam "Öğrenci" kelimesi geçen hücre: {student_found}')
