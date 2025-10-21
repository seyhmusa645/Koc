import docx
import json
import re

# Word dosyasını oku
doc = docx.Document('Sonuclar-Ortaokul-Rapor.docx')

print(f'Word dosyasında tablo sayısı: {len(doc.tables)}')

all_students = []

# Her tabloyu analiz et
for table_num, table in enumerate(doc.tables):
    student_data = {}
    
    # Tablonun tüm hücrelerini oku
    for row in table.rows:
        for cell in row.cells:
            cell_text = cell.text.strip()
            
            # Öğrenci No
            if 'Öğrenci No:' in cell_text:
                student_no = cell_text.replace('Öğrenci No:', '').strip()
                student_data['student_no'] = student_no
            
            # Adı Soyadı
            elif 'Adı Soyadı:' in cell_text:
                name = cell_text.replace('Adı Soyadı:', '').strip()
                student_data['name'] = name
            
            # Cinsiyeti
            elif 'Cinsiyeti:' in cell_text:
                gender = cell_text.replace('Cinsiyeti:', '').strip()
                student_data['gender'] = gender
            
            # Sınıfı
            elif 'Sınıfı:' in cell_text:
                grade_class = cell_text.replace('Sınıfı:', '').strip()
                student_data['grade_class'] = grade_class
                
                # Sınıf ve şube bilgisini ayır
                if '/' in grade_class:
                    grade, class_name = grade_class.split('/', 1)
                    student_data['grade'] = grade.strip()
                    student_data['class'] = class_name.strip()
                else:
                    student_data['grade'] = grade_class.strip()
                    student_data['class'] = None
            
            # Öğrenme Stili
            elif 'Öğrencinin Öğrenme Stili:' in cell_text:
                learning_style = cell_text.replace('Öğrencinin Öğrenme Stili:', '').strip()
                student_data['learning_style'] = learning_style
    
    # Eğer öğrenci verisi tamamsa ekle
    if 'name' in student_data and 'learning_style' in student_data:
        all_students.append(student_data)
        student_data = {}  # Yeni öğrenci için sıfırla

print(f'Toplam öğrenci bulundu: {len(all_students)}')

# İlk 10 öğrenciyi göster
print('\n=== İLK 10 ÖĞRENCİ ===')
for i, student in enumerate(all_students[:10]):
    print(f'{i+1}. {student.get("name", "Bilinmeyen")} - {student.get("grade_class", "Bilinmeyen")} - {student.get("learning_style", "Bilinmeyen")}')

# Öğrenme stilleri analizi
learning_styles = {}
for student in all_students:
    style = student.get('learning_style', 'Belirsiz')
    learning_styles[style] = learning_styles.get(style, 0) + 1

print('\n=== ÖĞRENME STİLLERİ DAĞILIMI ===')
for style, count in sorted(learning_styles.items(), key=lambda x: x[1], reverse=True):
    print(f'{style}: {count} öğrenci')

# Sınıf dağılımı
grades = {}
for student in all_students:
    grade = student.get('grade', 'Belirsiz')
    grades[grade] = grades.get(grade, 0) + 1

print('\n=== SINIF DAĞILIMI ===')
for grade, count in sorted(grades.items()):
    print(f'{grade}. Sınıf: {count} öğrenci')

# Veriyi kaydet
with open('word_students_data.json', 'w', encoding='utf-8') as f:
    json.dump(all_students, f, ensure_ascii=False, indent=2)

print(f'\nVeriler word_students_data.json dosyasına kaydedildi')
