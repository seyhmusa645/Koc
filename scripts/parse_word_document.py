import docx
import json
import re

# Word dosyasını oku
doc = docx.Document('Sonuclar-Ortaokul-Rapor.docx')

print(f'Word dosyası toplam paragraf sayısı: {len(doc.paragraphs)}')

# Tüm metni birleştir
full_text = ''
for paragraph in doc.paragraphs:
    full_text += paragraph.text + '\n'

print(f'Toplam karakter sayısı: {len(full_text)}')

# Metnin bir kısmını göster
print('\n=== METNİN İLK 1000 KARAKTERİ ===')
print(full_text[:1000])

# Farklı pattern'ler deneyelim
patterns = [
    r'Öğrenci No:\s*(\d+).*?Adı Soyadı:\s*([^Cinsiyeti]+)Cinsiyeti:\s*([^Sınıfı]+)Sınıfı:\s*([^\n]+).*?Öğrencinin Öğrenme Stili:\s*([^\n]+)',
    r'Adı Soyadı:\s*([^Cinsiyeti]+)Cinsiyeti:\s*([^Sınıfı]+)Sınıfı:\s*([^\n]+)',
    r'Öğrencinin Öğrenme Stili:\s*([^\n]+)',
    r'Adı Soyadı:\s*([^\n]+)',
    r'Öğrenci No:\s*(\d+)',
]

for i, pattern in enumerate(patterns):
    matches = re.findall(pattern, full_text, re.DOTALL)
    print(f'\nPattern {i+1}: {len(matches)} eşleşme')
    if matches and len(matches) > 0:
        print(f'İlk eşleşme: {matches[0]}')

# "Adı Soyadı" geçen yerleri bul
name_matches = re.findall(r'Adı Soyadı:\s*([^\n]+)', full_text)
print(f'\n"Adı Soyadı" geçen yerler: {len(name_matches)}')
for i, name in enumerate(name_matches[:10]):
    print(f'{i+1}. {name}')

# "Öğrenci No" geçen yerleri bul
student_no_matches = re.findall(r'Öğrenci No:\s*(\d+)', full_text)
print(f'\n"Öğrenci No" geçen yerler: {len(student_no_matches)}')
for i, no in enumerate(student_no_matches[:10]):
    print(f'{i+1}. {no}')

# "Öğrenme Stili" geçen yerleri bul
style_matches = re.findall(r'Öğrencinin Öğrenme Stili:\s*([^\n]+)', full_text)
print(f'\n"Öğrenme Stili" geçen yerler: {len(style_matches)}')
for i, style in enumerate(style_matches[:10]):
    print(f'{i+1}. {style}')
