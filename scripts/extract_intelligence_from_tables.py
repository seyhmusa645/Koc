import json
from docx import Document
import re

# Word dosyasını oku
doc = Document('Sonuclar-Ortaokul-Rapor.docx')

print("=== WORD DOSYASINDAKİ TABLOLAR ===")

# Tüm tabloları kontrol et
for table_idx, table in enumerate(doc.tables):
    print(f"\n--- TABLO {table_idx + 1} ---")
    print(f"Satır sayısı: {len(table.rows)}")
    print(f"Sütun sayısı: {len(table.columns)}")
    
    # İlk birkaç satırı göster
    for row_idx, row in enumerate(table.rows[:3]):
        row_data = []
        for cell in row.cells:
            row_data.append(cell.text.strip())
        print(f"Satır {row_idx + 1}: {row_data}")
    
    # Zeka türü içeren satırları ara
    for row_idx, row in enumerate(table.rows):
        row_text = ' '.join([cell.text.strip() for cell in row.cells])
        if any(zeka in row_text.upper() for zeka in ['ZEKA', 'SÖZEL', 'MATEMATİK', 'GÖRSEL', 'MÜZİK', 'BEDEN', 'KİŞİLER', 'İÇSEL', 'DOĞA']):
            print(f"ZEKA TÜRÜ BULUNDU - Satır {row_idx + 1}: {row_text}")

print("\n=== TÜM PARAGRAFLARDA ZEKA TÜRÜ ARAMA ===")
# Paragraflarda da ara
for para_idx, para in enumerate(doc.paragraphs):
    if any(zeka in para.text.upper() for zeka in ['ZEKA', 'SÖZEL', 'MATEMATİK', 'GÖRSEL', 'MÜZİK', 'BEDEN', 'KİŞİLER', 'İÇSEL', 'DOĞA']):
        print(f"Paragraf {para_idx + 1}: {para.text.strip()[:100]}...")
