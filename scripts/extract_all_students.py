import docx
import json
import re

# Word dosyasını oku
doc = docx.Document('Sonuclar-Ortaokul-Rapor.docx')

print(f'Word dosyasında tablo sayısı: {len(doc.tables)}')

all_students = []
current_student = {}

# Her tabloyu analiz et
for table_num, table in enumerate(doc.tables):
    # Tablonun tüm hücrelerini oku
    for row in table.rows:
        for cell in row.cells:
            cell_text = cell.text.strip()
            
            # Öğrenci No
            if 'Öğrenci No:' in cell_text:
                # Eğer önceki öğrenci varsa kaydet
                if current_student and 'name' in current_student:
                    all_students.append(current_student)
                
                # Yeni öğrenci başlat
                current_student = {}
                student_no = cell_text.replace('Öğrenci No:', '').strip()
                current_student['student_no'] = student_no
            
            # Adı Soyadı
            elif 'Adı Soyadı:' in cell_text:
                name = cell_text.replace('Adı Soyadı:', '').strip()
                current_student['name'] = name
            
            # Cinsiyeti
            elif 'Cinsiyeti:' in cell_text:
                gender = cell_text.replace('Cinsiyeti:', '').strip()
                current_student['gender'] = gender
            
            # Sınıfı
            elif 'Sınıfı:' in cell_text:
                grade_class = cell_text.replace('Sınıfı:', '').strip()
                current_student['grade_class'] = grade_class
                
                # Sınıf ve şube bilgisini ayır
                if '/' in grade_class:
                    grade, class_name = grade_class.split('/', 1)
                    current_student['grade'] = grade.strip()
                    current_student['class'] = class_name.strip()
                else:
                    current_student['grade'] = grade_class.strip()
                    current_student['class'] = None
            
            # Öğrenme Stili
            elif 'Öğrencinin Öğrenme Stili:' in cell_text:
                learning_style = cell_text.replace('Öğrencinin Öğrenme Stili:', '').strip()
                current_student['learning_style'] = learning_style

# Son öğrenciyi de ekle
if current_student and 'name' in current_student:
    all_students.append(current_student)

print(f'Toplam öğrenci bulundu: {len(all_students)}')

# İlk 10 öğrenciyi göster
print('\n=== İLK 10 ÖĞRENCİ ===')
for i, student in enumerate(all_students[:10]):
    print(f'{i+1}. {student.get("name", "Bilinmeyen")} - {student.get("grade_class", "Bilinmeyen")} - {student.get("learning_style", "Bilinmeyen")}')

# Öğrenme stilleri analizi
learning_styles = {}
for student in all_students:
    style = student.get('learning_style', 'Belirsiz')
    learning_styles[style] = learning_styles.get(style, 0) + 1

print('\n=== ÖĞRENME STİLLERİ DAĞILIMI ===')
for style, count in sorted(learning_styles.items(), key=lambda x: x[1], reverse=True):
    print(f'{style}: {count} öğrenci')

# Sınıf dağılımı
grades = {}
for student in all_students:
    grade = student.get('grade', 'Belirsiz')
    grades[grade] = grades.get(grade, 0) + 1

print('\n=== SINIF DAĞILIMI ===')
for grade, count in sorted(grades.items()):
    print(f'{grade}. Sınıf: {count} öğrenci')

# Cinsiyet dağılımı
genders = {}
for student in all_students:
    gender = student.get('gender', 'Belirsiz')
    genders[gender] = genders.get(gender, 0) + 1

print('\n=== CİNSİYET DAĞILIMI ===')
for gender, count in sorted(genders.items()):
    print(f'{gender}: {count} öğrenci')

# Veriyi kaydet
with open('complete_students_data.json', 'w', encoding='utf-8') as f:
    json.dump(all_students, f, ensure_ascii=False, indent=2)

print(f'\nVeriler complete_students_data.json dosyasına kaydedildi')
