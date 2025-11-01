#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Akıllı Karne PDF Çıkarım Sistemi v2.0
Çoklu format desteği, otomatik tespit, güven skoru ve doğrulama
"""

import re
import sys
import os
import fitz  # PyMuPDF

# Modülleri import et
sys.path.insert(0, os.path.dirname(__file__))

from modules.detector import FormatDetector
from modules.extractors.template_extractor import TemplateExtractor
from modules.extractors.table_extractor import TableExtractor
from modules.extractors.text_extractor import TextExtractor
from modules.extractors.ocr_extractor import OCRExtractor
from modules import validators
from modules.logger import get_logger
from modules.normalizers import clean_page_text


def extract_to_excel_v2(input_file, output_file, templates_dir="templates"):
    """
    PDF'den Excel'e akıllı çıkarım (v2.0)
    
    Args:
        input_file: PDF dosya yolu
        output_file: Çıktı Excel dosyası
        templates_dir: Şablon dizini
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("HATA: openpyxl kütüphanesi yüklü değil!")
        print("Yüklemek için: pip install openpyxl")
        return
    
    # Logger başlat
    logger = get_logger()
    logger.log_pdf_info(input_file, 0)
    
    # PDF'i aç
    if not os.path.exists(input_file):
        logger.log_error(f"PDF dosyası bulunamadı: {input_file}")
        print(f"HATA: {input_file} dosyası bulunamadı!")
        return
    
    try:
        doc = fitz.open(input_file)
        logger.log_pdf_info(input_file, len(doc))
        print(f"\n[OK] PDF acildi: {len(doc)} sayfa")
    except Exception as e:
        logger.log_error("PDF açılamadı", e)
        print(f"HATA: PDF açılamadı: {e}")
        return
    
    # 1. FORMAT TESPİTİ
    print("\n[1/5] Format tespit ediliyor...")
    detector = FormatDetector(templates_dir)
    template, confidence, all_scores = detector.detect_format(input_file)
    
    if template:
        logger.log_format_detection(template['name'], confidence, all_scores)
        print(f"[OK] Format: {template['name']} (Guven: {confidence:.0%})")
    else:
        logger.log_format_detection(None, confidence, all_scores)
        print(f"[UYARI] Format taninamadi! Genel cikarim denenecek...")
    
    # 2. ÇIKARICIları HAZıRLA
    print("\n[2/5] Çıkarıcılar hazırlanıyor...")
    extractors = [
        TemplateExtractor(),
        TextExtractor(),
        TableExtractor(),
        OCRExtractor()
    ]
    print(f"[OK] {len(extractors)} cikarici hazir")
    
    # 3. SAYFALARI İŞLE
    print(f"\n[3/5] Sayfalar işleniyor...")
    
    valid_records = []
    suspicious_records = []
    skipped_records = []
    
    # Sayfaları ayır ("SONUÇ BELGESİ" ile)
    all_text = ""
    for page_num in range(len(doc)):
        all_text += doc[page_num].get_text()
    
    pages = re.split(r'(?=SONUÇ BELGESİ)', all_text)
    total_pages = len(pages)
    
    print(f"  {total_pages} bölüm bulundu")
    
    for page_idx, page_text in enumerate(pages):
        page_num = page_idx + 1
        
        # Boş veya başlık sayfası kontrolü
        if not page_text.strip() or 'SONUÇ BELGESİ' not in page_text:
            continue
        
        logger.log_page_processing(page_num, total_pages)
        
        # Metni temizle
        page_text = clean_page_text(page_text)
        
        # 3.1. ÇIKARICIları SIRALA DENE
        best_data = {}
        best_confidence = 0.0
        best_extractor = None
        
        for extractor in extractors:
            # Şablon desteği kontrolü
            if not extractor.supports_template(template):
                continue
            
            logger.log_extractor_attempt(extractor.get_name(), page_num)
            
            try:
                data, conf = extractor.extract(page_text, template, None)
                
                if conf > best_confidence:
                    best_data = data
                    best_confidence = conf
                    best_extractor = extractor.get_name()
                
                logger.log_extractor_result(extractor.get_name(), page_num, 
                                           bool(data), conf, 
                                           validators.get_validation_summary(data) if data else None)
                
                # Yüksek güvenle bulunduysa dur
                if conf >= 0.85:
                    break
                    
            except Exception as e:
                logger.log_debug(f"  [{page_num}] {extractor.get_name()} hata: {e}")
                continue
        
        # 3.2. VERİYİ DOĞRULA
        if not best_data.get('ad_soyad'):
            reason = "Öğrenci bilgisi bulunamadı"
            logger.log_student_skipped(page_num, reason)
            skipped_records.append({'page': page_num, 'reason': reason})
            continue
        
        is_valid, val_confidence, issues = validators.validate_student_data(best_data, template)
        
        # Final güven skoru = çıkarım güveni * doğrulama güveni
        final_confidence = best_confidence * val_confidence
        
        # Kayıt ekle
        student_name = best_data.get('ad_soyad', 'Bilinmiyor')
        
        if final_confidence >= 0.6:  # %60 ve üzeri geçerli
            # Geçerli kayıt
            logger.log_validation_result(page_num, student_name, True, final_confidence, issues)
            best_data['_confidence'] = final_confidence
            best_data['_extractor'] = best_extractor
            valid_records.append(best_data)
            print(f"  [OK] [{page_num}] {student_name} - {best_extractor}")
        else:
            # Şüpheli kayıt
            logger.log_validation_result(page_num, student_name, False, final_confidence, issues)
            best_data['_confidence'] = final_confidence
            best_data['_extractor'] = best_extractor
            best_data['_issues'] = '; '.join(issues)
            suspicious_records.append(best_data)
            print(f"  [SPHELI] [{page_num}] {student_name} - Supheli ({final_confidence:.0%})")
    
    doc.close()
    
    # Özet
    logger.log_summary(total_pages, len(valid_records), len(suspicious_records), len(skipped_records))
    print(f"\n[4/5] İşlem tamamlandı")
    print(f"  Geçerli: {len(valid_records)}")
    print(f"  Şüpheli: {len(suspicious_records)}")
    print(f"  Atlanan: {len(skipped_records)}")
    
    # 4. EXCEL OLUŞTUR
    print(f"\n[5/5] Excel dosyası oluşturuluyor...")
    
    wb = openpyxl.Workbook()
    
    # Sheet 1: Geçerli Kayıtlar
    ws_valid = wb.active
    ws_valid.title = "Öğrenci Karneleri"
    
    # Başlıklar
    headers = [
        'Ad Soyad', 'Numara', 'Sınıf', 'Katılımlar',
        'LGS Puanı', 'LGS Ortalama', 'Sınıf Derecesi', 'Kurum Derecesi',
        'İlçe Derecesi', 'İl Derecesi', 'Genel Derece',
        'Türkçe Doğru', 'Türkçe Yanlış', 'Türkçe Net', 'Türkçe Başarı%',
        'Sosyal/İnkılap Doğru', 'Sosyal/İnkılap Yanlış', 'Sosyal/İnkılap Net', 'Sosyal/İnkılap Başarı%',
        'Din Doğru', 'Din Yanlış', 'Din Net', 'Din Başarı%',
        'İngilizce Doğru', 'İngilizce Yanlış', 'İngilizce Net', 'İngilizce Başarı%',
        'Matematik Doğru', 'Matematik Yanlış', 'Matematik Net', 'Matematik Başarı%',
        'Fen Doğru', 'Fen Yanlış', 'Fen Net', 'Fen Başarı%',
        'Toplam Doğru', 'Toplam Yanlış', 'Toplam Net', 'Toplam Başarı%',
        'Kazanım Detayları', 'Güven Skoru', 'Çıkarıcı'
    ]
    
    # Başlık stili
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for col, header in enumerate(headers, start=1):
        cell = ws_valid.cell(row=1, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Geçerli kayıtları yaz
    for row_idx, data in enumerate(valid_records, start=2):
        row_data = [
            data.get('ad_soyad', ''),
            data.get('numara', ''),
            data.get('sinif', ''),
            data.get('katilimlar', ''),
            data.get('lgs_puan', ''),
            data.get('lgs_ortalama', ''),
            data.get('sinif_derecesi', ''),
            data.get('kurum_derecesi', ''),
            data.get('ilce_derecesi', ''),
            data.get('il_derecesi', ''),
            data.get('genel_derece', ''),
            data.get('turk_dogru', ''),
            data.get('turk_yanlis', ''),
            data.get('turk_net', ''),
            data.get('turk_basari', ''),
            data.get('sosyal_dogru', ''),
            data.get('sosyal_yanlis', ''),
            data.get('sosyal_net', ''),
            data.get('sosyal_basari', ''),
            data.get('din_dogru', ''),
            data.get('din_yanlis', ''),
            data.get('din_net', ''),
            data.get('din_basari', ''),
            data.get('ing_dogru', ''),
            data.get('ing_yanlis', ''),
            data.get('ing_net', ''),
            data.get('ing_basari', ''),
            data.get('mat_dogru', ''),
            data.get('mat_yanlis', ''),
            data.get('mat_net', ''),
            data.get('mat_basari', ''),
            data.get('fen_dogru', ''),
            data.get('fen_yanlis', ''),
            data.get('fen_net', ''),
            data.get('fen_basari', ''),
            data.get('toplam_dogru', ''),
            data.get('toplam_yanlis', ''),
            data.get('toplam_net', ''),
            data.get('toplam_basari', ''),
            data.get('kazanimlar_ozet', 'Yok'),
            f"{data.get('_confidence', 0):.0%}",
            data.get('_extractor', '')
        ]
        
        for col, value in enumerate(row_data, start=1):
            ws_valid.cell(row=row_idx, column=col, value=value)
    
    # Sheet 2: Şüpheli Kayıtlar
    if suspicious_records:
        ws_suspicious = wb.create_sheet("İnceleme Gerekli")
        
        susp_headers = headers + ['Sorunlar']
        for col, header in enumerate(susp_headers, start=1):
            cell = ws_suspicious.cell(row=1, column=col)
            cell.value = header
            cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        for row_idx, data in enumerate(suspicious_records, start=2):
            row_data = [
                data.get('ad_soyad', ''),
                data.get('numara', ''),
                data.get('sinif', ''),
                data.get('katilimlar', ''),
                data.get('lgs_puan', ''),
                data.get('lgs_ortalama', ''),
                data.get('sinif_derecesi', ''),
                data.get('kurum_derecesi', ''),
                data.get('ilce_derecesi', ''),
                data.get('il_derecesi', ''),
                data.get('genel_derece', ''),
                data.get('turk_dogru', ''),
                data.get('turk_yanlis', ''),
                data.get('turk_net', ''),
                data.get('turk_basari', ''),
                data.get('sosyal_dogru', ''),
                data.get('sosyal_yanlis', ''),
                data.get('sosyal_net', ''),
                data.get('sosyal_basari', ''),
                data.get('din_dogru', ''),
                data.get('din_yanlis', ''),
                data.get('din_net', ''),
                data.get('din_basari', ''),
                data.get('ing_dogru', ''),
                data.get('ing_yanlis', ''),
                data.get('ing_net', ''),
                data.get('ing_basari', ''),
                data.get('mat_dogru', ''),
                data.get('mat_yanlis', ''),
                data.get('mat_net', ''),
                data.get('mat_basari', ''),
                data.get('fen_dogru', ''),
                data.get('fen_yanlis', ''),
                data.get('fen_net', ''),
                data.get('fen_basari', ''),
                data.get('toplam_dogru', ''),
                data.get('toplam_yanlis', ''),
                data.get('toplam_net', ''),
                data.get('toplam_basari', ''),
                data.get('kazanimlar_ozet', 'Yok'),
                f"{data.get('_confidence', 0):.0%}",
                data.get('_extractor', ''),
                data.get('_issues', '')
            ]
            
            for col, value in enumerate(row_data, start=1):
                cell = ws_suspicious.cell(row=row_idx, column=col, value=value)
                # Şüpheli satırları vurgula
                if col <= 3:
                    cell.fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
    
    # Sütun genişliklerini ayarla
    for ws in [ws_valid] + ([ws_suspicious] if suspicious_records else []):
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    # Dosyayı kaydet
    wb.save(output_file)
    print(f"[OK] Excel dosyasi: {output_file}")
    print(f"  - {len(valid_records)} gecerli kayit")
    if suspicious_records:
        print(f"  - {len(suspicious_records)} supheli kayit (ayri sheet'te)")


def main():
    """Ana fonksiyon"""
    print("=" * 80)
    print("AKILLI KARNE PDF ÇIKARIM SİSTEMİ v2.0")
    print("Çoklu Format Desteği | Otomatik Tespit | Güven Skoru")
    print("=" * 80)
    print()
    
    if len(sys.argv) < 2:
        print("Kullanım: python karne_excel_v2.py <pdf_dosyasi> [excel_dosyasi]")
        print()
        print("Örnek:")
        print("  python karne_excel_v2.py karne.pdf")
        print("  python karne_excel_v2.py karne.pdf sonuc.xlsx")
        return
    
    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else "ogrenci_karneleri.xlsx"
    
    if not output_file.endswith('.xlsx'):
        output_file += '.xlsx'
    
    if not os.path.exists(input_file):
        print(f"HATA: {input_file} dosyası bulunamadı!")
        return
    
    extract_to_excel_v2(input_file, output_file)
    
    print()
    print("=" * 80)
    print("[OK] Islem tamamlandi!")
    print("=" * 80)


if __name__ == "__main__":
    main()

