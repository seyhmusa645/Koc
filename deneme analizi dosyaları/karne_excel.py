#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Karne PDF'sinden her öğrencinin bilgilerini Excel'e çıkartan betik
"""

import re
import sys
import os

def parse_kazanimlar(text):
    """Ders bazlı kazanım satırlarını özetler."""
    analysis_start = text.find('DERSLERE GÖRE ANALİZ')
    if analysis_start == -1:
        analysis_start = text.find('KAZANIM')
        if analysis_start == -1:
            return []

    analysis_text = text[analysis_start:]
    section_pattern = re.compile(r'(Sözel|Sayısal) \(([A-ZÇĞİÖŞÜ]+)\)')
    matches = list(section_pattern.finditer(analysis_text))
    if not matches:
        return []

    def is_number(token: str) -> bool:
        return bool(re.fullmatch(r'-?\d+(?:[.,]\d+)?', token))

    kazanimlar = []
    for idx, match in enumerate(matches):
        start = match.start()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(analysis_text)
        block = analysis_text[start:end]

        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue

        if lines and (lines[0].startswith('Sözel') or lines[0].startswith('Sayısal')):
            lines.pop(0)
        if lines:
            lines.pop(0)  # ders adı
        for label in ('S', 'D', 'Y', 'B%'):
            if lines and lines[0].upper() == label:
                lines.pop(0)

        i = 0
        while i < len(lines):
            if not lines[i]:
                i += 1
                continue

            desc_parts = []
            while i < len(lines) and not is_number(lines[i]):
                desc_parts.append(lines[i])
                i += 1
            desc = ' '.join(desc_parts).strip()
            if not desc:
                continue

            if 'KAZANIM' in desc.upper() and 'BEL' in desc.upper():
                skipped = 0
                while i < len(lines) and is_number(lines[i]) and skipped < 4:
                    i += 1
                    skipped += 1
                continue

            numbers = []
            while i < len(lines) and is_number(lines[i]) and len(numbers) < 4:
                numbers.append(lines[i])
                i += 1

            if len(numbers) == 4:
                basari_value = numbers[3].replace(',', '.')
                prefix = match.group(2)
                kazanimlar.append(f"{prefix}: {desc[:200]} ({basari_value}%)")

    return kazanimlar


def parse_student_page(text):
    """Bir sayfadaki öğrenci bilgilerini parse eder"""
    student_data = {}
    
    # Öğrenci adı
    name_match = re.search(r'Öğrenci.*?Numara.*?Sınıf\s+(\S.*?)\s+(\d+)\s+([\d\-A-Z]+)', text, re.DOTALL)
    if name_match:
        student_data['ad_soyad'] = name_match.group(1).strip()
        student_data['numara'] = name_match.group(2).strip()
        student_data['sinif'] = name_match.group(3).strip()
    
    # LGS Puanı ve derece
    lgs_match = re.search(r'LGS\s+([\d,]+)\s+([\d,]+)\s+(\d+)\s+(\d+)\s+(\d+)\s+(\d+)\s+(\d+)', text)
    if lgs_match:
        student_data['lgs_puan'] = lgs_match.group(1).replace(',', '.')
        student_data['lgs_ortalama'] = lgs_match.group(2).replace(',', '.')
        student_data['sinif_derecesi'] = lgs_match.group(3)
        student_data['kurum_derecesi'] = lgs_match.group(4)
        student_data['ilce_derecesi'] = lgs_match.group(5)
        student_data['il_derecesi'] = lgs_match.group(6)
        student_data['genel_derece'] = lgs_match.group(7)
    
    # Katılımlar
    katilim_match = re.search(r'Katılımlar:\s+([\d\s]+)', text)
    if katilim_match:
        student_data['katilimlar'] = katilim_match.group(1).strip()
    
    # Ders bilgileri
    # Türkçe (15 veya 20 soru)
    turk_match = re.search(r'Türkçe\s+(15|20)\s+(\d+)\s+(\d+)\s+([\d,]+)\s+(\d+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)', text)
    if turk_match:
        student_data['turk_dogru'] = turk_match.group(2)
        student_data['turk_yanlis'] = turk_match.group(3)
        student_data['turk_net'] = turk_match.group(4).replace(',', '.')
        student_data['turk_basari'] = turk_match.group(5)
    
    # Tarih veya Sosyal Bilgiler
    tarih_match = re.search(r'Tarih\s+10\s+(\d+)\s+(\d+)\s+([\d,]+)\s+(\d+)', text)
    if tarih_match:
        student_data['tarih_dogru'] = tarih_match.group(1)
        student_data['tarih_yanlis'] = tarih_match.group(2)
        student_data['tarih_net'] = tarih_match.group(3).replace(',', '.')
        student_data['tarih_basari'] = tarih_match.group(4)
    else:
        # Sosyal Bilgiler dene (büyük veya küçük harf)
        sos_match = re.search(r'(?:Sosyal Bilgiler|SOSYAL BİLGİLER)\s+10\s+(\d+)\s+(\d+)\s+([\d,]+)\s+(\d+)', text)
        if sos_match:
            student_data['tarih_dogru'] = sos_match.group(1)
            student_data['tarih_yanlis'] = sos_match.group(2)
            student_data['tarih_net'] = sos_match.group(3).replace(',', '.')
            student_data['tarih_basari'] = sos_match.group(4)
    
    # Din Kültürü
    din_match = re.search(r'Din K\.ve A\.B\.\s+10\s+(\d+)\s+(\d+)\s+([\d,]+)\s+(\d+)', text)
    if din_match:
        student_data['din_dogru'] = din_match.group(1)
        student_data['din_yanlis'] = din_match.group(2)
        student_data['din_net'] = din_match.group(3).replace(',', '.')
        student_data['din_basari'] = din_match.group(4)
    
    # İngilizce
    ing_match = re.search(r'İngilizce\s+10\s+(\d+)\s+(\d+)\s+([\d,]+)\s+(\d+)', text)
    if ing_match:
        student_data['ing_dogru'] = ing_match.group(1)
        student_data['ing_yanlis'] = ing_match.group(2)
        student_data['ing_net'] = ing_match.group(3).replace(',', '.')
        student_data['ing_basari'] = ing_match.group(4)
    
    # Matematik (15 veya 20 soru)
    mat_match = re.search(r'Matematik\s+(15|20)\s+(\d+)\s+(\d+)\s+([\d,]+)\s+(\d+)', text)
    if mat_match:
        student_data['mat_dogru'] = mat_match.group(2)
        student_data['mat_yanlis'] = mat_match.group(3)
        student_data['mat_net'] = mat_match.group(4).replace(',', '.')
        student_data['mat_basari'] = mat_match.group(5)
    
    # Fen (15 veya 20 soru)
    fen_match = re.search(r'Fen\s+(15|20)\s+(\d+)\s+(\d+)\s+([\d,]+)\s+(\d+)', text)
    if fen_match:
        student_data['fen_dogru'] = fen_match.group(2)
        student_data['fen_yanlis'] = fen_match.group(3)
        student_data['fen_net'] = fen_match.group(4).replace(',', '.')
        student_data['fen_basari'] = fen_match.group(5)
    
    # Toplam (75 veya 90 soru)
    toplam_match = re.search(r'Toplam:\s+(75|90)\s+(\d+)\s+(\d+)\s+([\d,]+)\s+(\d+)', text)
    if toplam_match:
        student_data['toplam_dogru'] = toplam_match.group(2)
        student_data['toplam_yanlis'] = toplam_match.group(3)
        student_data['toplam_net'] = toplam_match.group(4).replace(',', '.')
        student_data['toplam_basari'] = toplam_match.group(5)
    # Kazanım bilgileri - Excel için özet formatında
    kazanimlar_list = parse_kazanimlar(text)
    student_data['kazanimlar_ozet'] = ' | '.join(kazanimlar_list) if kazanimlar_list else 'Yok'

    return student_data

def extract_to_excel(input_file, output_file):
    """PDF'den Excel'e çıkarma"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("HATA: openpyxl kütüphanesi yüklü değil!")
        print("Yüklemek için: pip install openpyxl")
        return
    
    # Dosyayı oku
    if input_file.lower().endswith('.pdf'):
        try:
            import fitz
            doc = fitz.open(input_file)
            all_text = ""
            for page_num in range(len(doc)):
                page = doc[page_num]
                all_text += page.get_text()
            doc.close()
            content = all_text
        except ImportError:
            print("HATA: PyMuPDF (fitz) yüklü değil!")
            print("Yüklemek için: pip install PyMuPDF")
            return
    else:
        with open(input_file, 'r', encoding='utf-8') as f:
            content = f.read()
    
    # Sayfaları ayır
    pages = re.split(r'(?=SONUÇ BELGESİ)', content)
    
    # Excel workbook oluştur
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Öğrenci Karneleri"
    
    # Başlık satırı
    headers = [
        'Ad Soyad', 'Numara', 'Sınıf', 'Katılımlar',
        'LGS Puanı', 'LGS Ortalama', 'Sınıf Derecesi', 'Kurum Derecesi', 
        'İlçe Derecesi', 'İl Derecesi', 'Genel Derece',
        'Türkçe Doğru', 'Türkçe Yanlış', 'Türkçe Net', 'Türkçe Başarı%',
        'Tarih Doğru', 'Tarih Yanlış', 'Tarih Net', 'Tarih Başarı%',
        'Din Doğru', 'Din Yanlış', 'Din Net', 'Din Başarı%',
        'İngilizce Doğru', 'İngilizce Yanlış', 'İngilizce Net', 'İngilizce Başarı%',
        'Matematik Doğru', 'Matematik Yanlış', 'Matematik Net', 'Matematik Başarı%',
        'Fen Doğru', 'Fen Yanlış', 'Fen Net', 'Fen Başarı%',
        'Toplam Doğru', 'Toplam Yanlış', 'Toplam Net', 'Toplam Başarı%',
        'Kazanım Detayları'
    ]
    
    # Başlık stilini ayarla
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Öğrenci verilerini işle
    row = 2
    for page in pages:
        if not page.strip() or 'SONUÇ BELGESİ' not in page:
            continue
        
        student_data = parse_student_page(page)
        
        if not student_data.get('ad_soyad'):
            continue
        
        # Verileri yaz
        data_row = [
            student_data.get('ad_soyad', ''),
            student_data.get('numara', ''),
            student_data.get('sinif', ''),
            student_data.get('katilimlar', ''),
            student_data.get('lgs_puan', ''),
            student_data.get('lgs_ortalama', ''),
            student_data.get('sinif_derecesi', ''),
            student_data.get('kurum_derecesi', ''),
            student_data.get('ilce_derecesi', ''),
            student_data.get('il_derecesi', ''),
            student_data.get('genel_derece', ''),
            student_data.get('turk_dogru', ''),
            student_data.get('turk_yanlis', ''),
            student_data.get('turk_net', ''),
            student_data.get('turk_basari', ''),
            student_data.get('tarih_dogru', ''),
            student_data.get('tarih_yanlis', ''),
            student_data.get('tarih_net', ''),
            student_data.get('tarih_basari', ''),
            student_data.get('din_dogru', ''),
            student_data.get('din_yanlis', ''),
            student_data.get('din_net', ''),
            student_data.get('din_basari', ''),
            student_data.get('ing_dogru', ''),
            student_data.get('ing_yanlis', ''),
            student_data.get('ing_net', ''),
            student_data.get('ing_basari', ''),
            student_data.get('mat_dogru', ''),
            student_data.get('mat_yanlis', ''),
            student_data.get('mat_net', ''),
            student_data.get('mat_basari', ''),
            student_data.get('fen_dogru', ''),
            student_data.get('fen_yanlis', ''),
            student_data.get('fen_net', ''),
            student_data.get('fen_basari', ''),
            student_data.get('toplam_dogru', ''),
            student_data.get('toplam_yanlis', ''),
            student_data.get('toplam_net', ''),
            student_data.get('toplam_basari', ''),
            student_data.get('kazanimlar_ozet', 'Yok')
        ]
        
        for col, value in enumerate(data_row, start=1):
            ws.cell(row=row, column=col, value=value)
        
        row += 1
    
    # Sütun genişliklerini ayarla
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 150)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Dosyayı kaydet
    wb.save(output_file)
    print(f'\n[OK] Excel dosyası oluşturuldu: {output_file}')
    print(f"  Toplam {row-2} öğrenci kaydedildi.")

def main():
    """Ana fonksiyon"""
    print("=" * 80)
    print("KARNE PDF'DEN EXCEL'E ÇIKARMA ARACI")
    print("=" * 80)
    print()
    
    if len(sys.argv) < 2:
        print("Kullanım: python karne_excel.py <pdf_dosyasi> [excel_dosyasi]")
        print()
        print("Örnek:")
        print("  python karne_excel.py karne.pdf")
        print("  python karne_excel.py karne.pdf sonuc.xlsx")
        return
    
    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else "ogrenci_karneleri.xlsx"
    
    if not output_file.endswith('.xlsx'):
        output_file += '.xlsx'
    
    if not os.path.exists(input_file):
        print(f"HATA: {input_file} dosyası bulunamadı!")
        return
    
    extract_to_excel(input_file, output_file)
    
    print()
    print("=" * 80)
    print("İşlem tamamlandı!")
    print("=" * 80)

if __name__ == "__main__":
    main()
